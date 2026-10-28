# Gera exemplos/modelo-cadastro-produtos.xlsx — o modelo baixado no botão
# "Modelo" da tela Produtos (produtos.html → Planilha).
#
# Formato lido pelo Upload (produtos.js → _prodUploadParsearExcelLote): uma
# linha de cabeçalho (procurada nas 5 primeiras linhas, precisa ter SKU e Nome)
# e um produto por linha abaixo dela. Os nomes de coluna vêm de
# PROD_MODELO_CAMPOS em produtos.js (lidos daqui, para o modelo acompanhar o
# sistema); "*" é ignorado na leitura. Título, instrução e a faixa de seções
# ficam acima do cabeçalho e não são importados.
#
# Uso: python3 scripts/gerar-modelo-produtos.py   (precisa de: pip install openpyxl)
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, AZUL, AZUL_CLARO, CINZA, AMARELO = '1E3A8A', '4776EC', 'EEF3FF', '64748B', 'FFF8DB'
fino = Side(style='thin', color='CBD5E1')
borda = Border(left=fino, right=fino, top=fino, bottom=fino)
LINHAS_PREENCHER = 300

# ── campos de produtos.js ──
js = open('produtos.js', encoding='utf-8').read()
bloco = js[js.index('const PROD_MODELO_CAMPOS = ['):]
bloco = bloco[:bloco.index('\n];')]
def campo_js(obj, chave):
    m = re.search(chave + r":\s*(true|false|'((?:[^'\\]|\\.)*)')", obj)
    if not m:
        return ''
    return m.group(1) == 'true' if m.group(2) is None else m.group(2).replace("\\'", "'")
CAMPOS = [{k: campo_js(o, k) for k in ('secao', 'coluna', 'campo', 'obrigatorio', 'tipo', 'exemplo', 'obs')}
          for o in re.findall(r'\{[^{}]*coluna:[^{}]*\}', bloco)]
assert CAMPOS, 'PROD_MODELO_CAMPOS não encontrado em produtos.js'

TEXTO_LITERAL = {'sku', 'ncm', 'cest', 'gtin', 'naladi_nesh', 'dun14'}   # sem perder zeros
LARGURA = {'nome': 34, 'descricao': 40, 'obs_preco': 30, 'obs_estoque': 30, 'obs_logistica': 30, 'imagem_url': 30, 'empresa_parceira_ref': 28}

wb = Workbook()
ws = wb.active
ws.title = 'Produtos'
ws.sheet_view.showGridLines = False
ultima = get_column_letter(len(CAMPOS))

ws.merge_cells(f'A1:{ultima}1')
ws['A1'] = 'MODELO DE CADASTRO DE PRODUTOS'
ws['A1'].font = Font(bold=True, size=15, color='FFFFFF')
ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
ws['A1'].alignment = Alignment(vertical='center', indent=1)
ws.row_dimensions[1].height = 30
ws.merge_cells(f'A2:{ultima}2')
ws['A2'] = ('Um produto por linha, a partir da linha 5 (em amarelo). Campos com * são obrigatórios. '
            'Veja as abas "Exemplo" e "Instruções". Depois envie pelo botão Upload na tela Produtos.')
ws['A2'].font = Font(italic=True, size=10, color=CINZA)
ws['A2'].alignment = Alignment(vertical='center', indent=1)
ws.row_dimensions[2].height = 22

# faixa de seções (linha 3) e cabeçalho (linha 4)
inicio = 1
for i, c in enumerate(CAMPOS + [{'secao': None}], start=1):
    if i > len(CAMPOS) or c['secao'] != CAMPOS[inicio - 1]['secao']:
        a, b = get_column_letter(inicio), get_column_letter(i - 1)
        if a != b:
            ws.merge_cells(f'{a}3:{b}3')
        cel = ws[f'{a}3']
        cel.value = CAMPOS[inicio - 1]['secao'].upper()
        cel.font = Font(bold=True, size=9, color=NAVY)
        cel.alignment = Alignment(horizontal='center', vertical='center')
        for col in range(inicio, i):
            ws.cell(3, col).fill = PatternFill('solid', fgColor=AZUL_CLARO)
            ws.cell(3, col).border = borda
        inicio = i
ws.row_dimensions[3].height = 18

for col, c in enumerate(CAMPOS, start=1):
    cel = ws.cell(4, col, c['coluna'] + (' *' if c['obrigatorio'] else ''))
    cel.font = Font(bold=True, color='FFFFFF', size=10)
    cel.fill = PatternFill('solid', fgColor=AZUL)
    cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cel.border = borda
    ws.column_dimensions[get_column_letter(col)].width = LARGURA.get(c['campo'], max(13, len(c['coluna']) + 3))
ws.row_dimensions[4].height = 32
ws.freeze_panes = 'C5'   # cabeçalho e SKU/Nome sempre visíveis

# linhas de preenchimento
amarelo = PatternFill('solid', fgColor=AMARELO)
for lin in range(5, 5 + LINHAS_PREENCHER):
    for col, c in enumerate(CAMPOS, start=1):
        cel = ws.cell(lin, col)
        cel.fill = amarelo
        cel.border = borda
        if c['campo'] in TEXTO_LITERAL:
            cel.number_format = '@'

def validacao(campo, lista, aviso, estrita=True):
    col = next((i for i, c in enumerate(CAMPOS, start=1) if c['campo'] == campo), None)
    if not col:
        return
    dv = DataValidation(type='list', formula1=f'"{lista}"', allow_blank=True, showErrorMessage=estrita,
                        error=aviso, errorTitle='Valor inválido')
    ws.add_data_validation(dv)
    letra = get_column_letter(col)
    dv.add(f'{letra}5:{letra}{4 + LINHAS_PREENCHER}')

validacao('status', 'ativo,pendente,pausado,inativo', 'Escolha: ativo, pendente, pausado ou inativo.')
validacao('controla_estoque', 'Sim,Não', 'Escolha Sim ou Não.')
validacao('venda_sem_estoque', 'Sim,Não', 'Escolha Sim ou Não.')
validacao('idioma_nome', 'Português,Inglês,Espanhol,Francês,Alemão,Chinês', '', estrita=False)   # outro idioma também vale

# ── aba Exemplo ──
we = wb.create_sheet('Exemplo')
we.sheet_view.showGridLines = False
we['A1'] = 'EXEMPLO DE PRODUTO PREENCHIDO (esta aba não é importada)'
we['A1'].font = Font(bold=True, size=12, color=NAVY)
for col, c in enumerate(CAMPOS, start=1):
    h = we.cell(3, col, c['coluna'])
    h.font = Font(bold=True, color='FFFFFF', size=10)
    h.fill = PatternFill('solid', fgColor=AZUL)
    h.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    h.border = borda
    v = we.cell(4, col, c['exemplo'])
    v.font = Font(italic=True, color='334155')
    v.border = borda
    we.column_dimensions[get_column_letter(col)].width = ws.column_dimensions[get_column_letter(col)].width
we.row_dimensions[3].height = 32

# ── aba Instruções ──
wi = wb.create_sheet('Instruções')
wi.sheet_view.showGridLines = False
for col, larg in zip('ABCDEF', (16, 26, 12, 10, 30, 80)):
    wi.column_dimensions[col].width = larg
wi['A1'] = 'COMO PREENCHER'
wi['A1'].font = Font(bold=True, size=13, color=NAVY)
passos = [
    '1. Preencha a aba "Produtos": um produto por linha, a partir da linha 5. Não altere o cabeçalho (linha 4).',
    '2. SKU e Nome são obrigatórios — linhas sem eles são ignoradas. O SKU não pode se repetir.',
    '3. Salve o arquivo (.xlsx) e, na tela Produtos, clique em Upload e escolha este arquivo.',
    '4. Os textos são gravados em MAIÚSCULAS, como no formulário (tags, referências, lote e link da imagem ficam como digitados).',
]
for i, p in enumerate(passos, start=2):
    wi.merge_cells(f'A{i}:F{i}')
    wi[f'A{i}'] = p
    wi[f'A{i}'].font = Font(size=10, color='334155')
linha = len(passos) + 3
for col, txt in enumerate(('Seção', 'Coluna', 'Obrigatório', 'Tipo', 'Exemplo', 'Orientação'), start=1):
    c = wi.cell(linha, col, txt)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.border = borda
for c in CAMPOS:
    linha += 1
    for col, val in enumerate((c['secao'], c['coluna'], 'Sim' if c['obrigatorio'] else '', c['tipo'], c['exemplo'], c['obs']), start=1):
        cel = wi.cell(linha, col, val)
        cel.border = borda
        cel.alignment = Alignment(vertical='center', wrap_text=(col == 6))
        cel.font = Font(size=9, bold=(col == 2), color='1E293B' if col == 2 else '475569')
    wi.row_dimensions[linha].height = 28 if len(c['obs']) > 85 else 18

wb.active = 0
wb.save('exemplos/modelo-cadastro-produtos.xlsx')
print(f'gerado: exemplos/modelo-cadastro-produtos.xlsx ({len(CAMPOS)} colunas)')
