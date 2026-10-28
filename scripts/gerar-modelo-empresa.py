# Gera exemplos/modelo-cadastro-empresa.xlsx — o modelo baixado pelo botão
# "Modelo" da tela Empresas (cadastros.html).
#
# Formato lido pelo Upload (cadastros.js → _uploadParsearExcel): coluna A =
# nome do campo, coluna B = valor. Os nomes da coluna A precisam bater com os
# ALIASES de lá ("*" e o que estiver entre parênteses são ignorados na leitura).
# Linhas de título/seção e as colunas C e D (exemplo e orientação) são ignoradas.
#
# Uso: python3 scripts/gerar-modelo-empresa.py   (precisa de: pip install openpyxl)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, AZUL, AZUL_CLARO, CINZA, AMARELO = '1E3A8A', '4776EC', 'EEF3FF', '64748B', 'FFF8DB'
fino = Side(style='thin', color='CBD5E1')
borda = Border(left=fino, right=fino, top=fino, bottom=fino)

# (campo, obrigatório, exemplo, orientação)
SECOES = [
    ('CLASSIFICAÇÃO', [
        ('Modelo', False, 'Fabricante / Fornecedor', 'Escolha na lista. Em branco: Fabricante/Fornecedor (Brasil) ou Company (outro país).'),
        ('Tipo', True, 'Fabricante, Comprador', 'Um ou mais, separados por vírgula: Fabricante, Fornecedor, Comprador, Importador. (Transportadora não usa.)'),
        ('Código Interno', False, 'CLI-0001', 'Código da empresa no seu controle (aparece como "Código Interno" no Processo).'),
    ]),
    ('IDENTIFICAÇÃO', [
        ('Razão Social', True, 'EXEMPLO COMERCIAL LTDA', 'Nome jurídico completo da empresa.'),
        ('Nome Fantasia', False, 'EXEMPLO', 'Nome comercial, se houver.'),
        ('CNPJ', False, '12.345.678/0001-90', 'Empresa brasileira. Pessoa física: deixe em branco e use a linha CPF.'),
        ('CPF', False, '123.456.789-00', 'Só para pessoa física.'),
        ('Inscrição Estadual', False, '123.456.789.110', 'Obrigatória para empresa brasileira. Se não tiver, escreva ISENTO.'),
        ('Inscrição Municipal', False, '1.234.567-8', 'Opcional (Fabricante/Fornecedor e Trading).'),
        ('SUFRAMA', False, '', 'Só para empresas da Zona Franca de Manaus.'),
        ('País', False, 'Brasil', 'Em branco = Brasil. Outro país = empresa estrangeira (o nº fiscal — TIN, VAT, EIN — vai no formulário).'),
    ]),
    ('ENDEREÇO', [
        ('CEP', False, '01001-000', 'Empresa brasileira: com o CEP, o endereço é completado automaticamente.'),
        ('Estado', False, 'SP', 'Sigla do estado (UF) — escolha na lista.'),
        ('Cidade', False, 'São Paulo', ''),
        ('Bairro', False, 'Centro', ''),
        ('Endereço', False, 'Praça da Sé', 'Rua, avenida, rodovia...'),
        ('Número', False, '100', 'Sem número: escreva S/N.'),
        ('Complemento', False, 'Sala 1', 'Opcional.'),
    ]),
    ('CONTATO PRINCIPAL', [
        ('Nome do Contato', False, 'Maria Souza', 'Pessoa de contato na empresa.'),
        ('Cargo', False, 'Financeiro', 'Cargo ou área do contato (Comercial, Financeiro, Logística...).'),
        ('E-mail', False, 'contato@exemplo.com.br', 'E-mail do contato principal.'),
        ('Telefone', False, '(11) 99999-9999', 'Telefone do contato principal.'),
    ]),
    ('COMPLEMENTOS', [
        ('Site', False, 'www.exemplo.com.br', ''),
        ('Horário de Atendimento', False, '08:00 às 18:00', ''),
        ('Rede Social', False, '@exemplo', 'Instagram, LinkedIn...'),
        ('Tags', False, 'cliente vip; têxtil', 'Até 4 palavras-chave, separadas por ponto e vírgula.'),
        ('Observações', False, 'Prefere contato por e-mail', 'Opcional.'),
    ]),
]
MODELOS = 'Fabricante / Fornecedor,Company (estrangeira),Transportadora,Trading,Outros'
UFS = 'AC,AL,AP,AM,BA,CE,DF,ES,GO,MA,MT,MS,MG,PA,PB,PR,PE,PI,RJ,RN,RS,RO,RR,SC,SP,SE,TO'

wb = Workbook()
ws = wb.active
ws.title = 'Cadastro de Empresa'
ws.sheet_view.showGridLines = False
for col, larg in zip('ABCD', (24, 40, 28, 70)):
    ws.column_dimensions[col].width = larg

# Título e instrução
ws.merge_cells('A1:D1')
ws['A1'] = 'MODELO DE CADASTRO DE EMPRESA'
ws['A1'].font = Font(bold=True, size=15, color='FFFFFF')
ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
ws['A1'].alignment = Alignment(vertical='center', indent=1)
ws.row_dimensions[1].height = 30
ws.merge_cells('A2:D2')
ws['A2'] = 'Preencha a coluna VALOR (em amarelo) e envie pelo botão Upload na tela Empresas — 1 empresa por arquivo. Campos com * são obrigatórios.'
ws['A2'].font = Font(italic=True, size=10, color=CINZA)
ws['A2'].alignment = Alignment(vertical='center', indent=1, wrap_text=True)
ws.row_dimensions[2].height = 30

# Cabeçalho da tabela
linha = 4
for col, txt in zip('ABCD', ('Campo', 'Valor', 'Exemplo', 'Orientação')):
    c = ws[f'{col}{linha}']
    c.value = txt
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.alignment = Alignment(vertical='center', indent=1)
    c.border = borda
ws.row_dimensions[linha].height = 22
ws.freeze_panes = f'A{linha + 1}'

dv_uf = DataValidation(type='list', formula1=f'"{UFS}"', allow_blank=True,
                       error='Escolha a sigla do estado na lista.', errorTitle='Estado')
ws.add_data_validation(dv_uf)
dv_modelo = DataValidation(type='list', formula1=f'"{MODELOS}"', allow_blank=True,
                           error='Escolha o modelo na lista.', errorTitle='Modelo')
ws.add_data_validation(dv_modelo)

for secao, campos in SECOES:
    linha += 1
    ws.merge_cells(f'A{linha}:D{linha}')
    c = ws[f'A{linha}']
    c.value = secao
    c.font = Font(bold=True, size=10, color=NAVY)
    c.fill = PatternFill('solid', fgColor=AZUL_CLARO)
    c.alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[linha].height = 20
    for campo, obrigatorio, exemplo, orientacao in campos:
        linha += 1
        a, b, cc, d = (ws[f'{x}{linha}'] for x in 'ABCD')
        a.value = campo + (' *' if obrigatorio else '')
        a.font = Font(bold=True, color='1E293B')
        b.fill = PatternFill('solid', fgColor=AMARELO)
        b.number_format = '@'   # texto: CNPJ, CEP e IE não perdem zeros/formatação
        cc.value = exemplo
        cc.font = Font(italic=True, color=CINZA)
        d.value = orientacao
        d.font = Font(size=9, color=CINZA)
        for x in (a, b, cc, d):
            x.border = borda
            x.alignment = Alignment(vertical='center', indent=1, wrap_text=(x is d))
        ws.row_dimensions[linha].height = 30 if len(orientacao) > 70 else 20
        if campo == 'Estado':
            dv_uf.add(b)
        if campo == 'Modelo':
            dv_modelo.add(b)

ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

# Aba de instruções
wi = wb.create_sheet('Instruções')
wi.sheet_view.showGridLines = False
wi.column_dimensions['A'].width = 100
passos = [
    ('COMO USAR ESTE MODELO', True),
    ('1. Preencha a coluna "Valor" (em amarelo) da aba "Cadastro de Empresa". A coluna "Exemplo" é só referência.', False),
    ('2. Não altere a coluna "Campo" — é ela que o sistema usa para reconhecer cada informação.', False),
    ('3. Salve o arquivo (.xlsx) e, na tela Empresas, clique em Upload e escolha este arquivo.', False),
    ('4. O formulário de cadastro abre já preenchido: confira, complete o que faltar (marcas, contas bancárias, documentos...) e salve.', False),
    ('', False),
    ('Dicas', True),
    ('• Um arquivo = uma empresa. Para cadastrar outra, use uma cópia nova do modelo.', False),
    ('• Empresa brasileira: com o CEP preenchido, o endereço é completado automaticamente no formulário.', False),
    ('• Empresa estrangeira: informe o País; o número fiscal (TIN, VAT, EIN) é preenchido no formulário.', False),
    ('• Os textos são gravados em MAIÚSCULAS, como no restante do sistema (o e-mail fica em minúsculas).', False),
]
for i, (txt, titulo) in enumerate(passos, start=1):
    c = wi[f'A{i}']
    c.value = txt
    c.font = Font(bold=titulo, size=13 if i == 1 else 11, color=NAVY if titulo else '334155')
    c.alignment = Alignment(wrap_text=True, vertical='center')
    wi.row_dimensions[i].height = 24

wb.active = 0
wb.save('exemplos/modelo-cadastro-empresa.xlsx')
print('gerado: exemplos/modelo-cadastro-empresa.xlsx')
